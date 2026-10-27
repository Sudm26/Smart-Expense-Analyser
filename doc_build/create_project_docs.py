from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)

NAVY = "17243F"
BLUE = "315F9D"
PALE_BLUE = "EEF4FA"
LIGHT_GRAY = "D9D9D9"
MUTED = "5E6A7A"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_border(cell, color=LIGHT_GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def set_font(run, size=10.5, bold=False, color="000000", italic=False):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for style_name, size in (("Title", 23), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(13 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Smart Expense and Receipt Analyzer")
    set_font(run, size=8, color=MUTED)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        set_font(run, bold=True)
        run = p.add_run(text[len(bold_prefix):])
        set_font(run)
    else:
        run = p.add_run(text)
        set_font(run)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(text))
    return p


def make_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for idx, value in enumerate(headers):
        cell = header_cells[idx]
        if widths:
            cell.width = Inches(widths[idx])
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(value), size=9, bold=True, color="FFFFFF")
    for row_idx, row_values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cell = cells[idx]
            if widths:
                cell.width = Inches(widths[idx])
            if row_idx % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(str(value)), size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def draw_bar_chart(path, title, labels, values, suffix="%"):
    width, height = 1100, 540
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("arial.ttf", 28)
    small = ImageFont.truetype("arial.ttf", 21)
    bold = ImageFont.truetype("arialbd.ttf", 24)
    draw.text((55, 32), title, fill="#17243F", font=bold)
    x0, y0, x1, y1 = 130, 110, 1030, 440
    draw.line((x0, y0, x0, y1), fill="#687386", width=2)
    draw.line((x0, y1, x1, y1), fill="#687386", width=2)
    maximum = max(max(values), 100)
    for percent in range(0, 101, 25):
        y = y1 - (percent / maximum) * (y1 - y0)
        draw.line((x0, y, x1, y), fill="#E5EAF0", width=1)
        draw.text((55, y - 12), f"{percent}{suffix}", fill="#5E6A7A", font=small)
    bar_width = 160
    gap = (x1 - x0 - len(values) * bar_width) / (len(values) + 1)
    colors = ["#315F9D", "#3D9D7A", "#6D5DFC", "#FFB547"]
    for i, (label, value) in enumerate(zip(labels, values)):
        left = x0 + gap + i * (bar_width + gap)
        top = y1 - (value / maximum) * (y1 - y0)
        draw.rounded_rectangle((left, top, left + bar_width, y1), radius=8, fill=colors[i % len(colors)])
        draw.text((left + 42, top - 35), f"{value}{suffix}", fill="#17243F", font=bold)
        words = label.split()
        draw.text((left + 10, y1 + 18), "\n".join(words), fill="#17243F", font=small, align="center")
    draw.text((130, 495), "Preliminary functional checks; not a full held-out benchmark.", fill="#5E6A7A", font=small)
    image.save(path)


def project_report():
    doc = Document()
    configure(doc)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Smart Expense and Receipt Analyzer"), size=23, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("AI Project Report"), size=12, color=MUTED)
    doc.add_paragraph()
    make_table(doc, ["Item", "Project statement"], [
        ["Main AI technique", "Hybrid text classification using keyword rules, user-learned rules, TF-IDF features, and Logistic Regression."],
        ["Problem", "Manual expense tracking is slow and receipt information is difficult to organize reliably."],
        ["Dataset", "Project expense descriptions plus a local receipt-image collection used for OCR review and extraction testing."],
    ], [1.65, 5.25])

    heading(doc, "Abstract")
    para(doc, "The Smart Expense and Receipt Analyzer is a web application that converts receipts and short manual descriptions into organized expense records. Users can upload a receipt image or add an expense without a receipt, such as snacks, stationery, or an autorickshaw ride. For receipt images, the system applies image preprocessing and EasyOCR in three variants: original, contrast-enhanced, and thresholded. It then uses layout-aware rules to identify a merchant, date, and final total while avoiding subtotal, tax, tip, tender, and item-count rows. For manual entries, a hybrid classifier combines verified keywords with TF-IDF and Logistic Regression. The application also supports user feedback: a corrected manual description can be remembered as a private category rule for future predictions. Spending records are stored in SQLite, summarized by category, and checked against user-defined monthly category thresholds. The project demonstrates practical AI through text classification, rule-based knowledge representation, supervised learning, and human-in-the-loop correction. The current evaluation is a functional prototype evaluation, not a large benchmark. Clear receipts and known descriptions perform best; ambiguous or faded text remains available for user review before saving. The system therefore prioritizes useful assistance and transparent correction over unsupported accuracy claims.")
    para(doc, "Keywords: receipt OCR; text classification; expert system; Logistic Regression; expense analytics", bold_prefix="Keywords: ")

    heading(doc, "Background")
    para(doc, "Students and households often lose receipts or postpone recording small daily purchases. Existing manual spreadsheets require repeated data entry and do not easily identify overspending. This project applies AI-assisted extraction and classification to make expense logging faster while retaining user control.")

    heading(doc, "Objective")
    bullet(doc, "Extract merchant, date, and final amount from receipt images.")
    bullet(doc, "Classify manual and receipt-based expenses into useful categories.")
    bullet(doc, "Learn approved manual corrections as private reusable rules.")
    bullet(doc, "Compare current-month spending against user-defined category thresholds.")

    heading(doc, "Proposed Method")
    para(doc, "The system receives either a receipt image or a manual description. Receipt images pass through preprocessing and OCR. Structured rules select the most likely merchant and final total from text and layout evidence. Manual descriptions first use user-learned and keyword rules, then fall back to a TF-IDF and Logistic Regression classifier. Every prediction remains editable before storage.")

    heading(doc, "Overall Architecture and Features")
    make_table(doc, ["Stage", "Implementation", "Output"], [
        ["Input", "Receipt upload or manual description", "Image or short text"],
        ["Processing", "OpenCV preprocessing and EasyOCR; text normalization", "OCR blocks and clean text"],
        ["AI decision", "Rules plus TF-IDF Logistic Regression", "Category and confidence"],
        ["Review", "Editable prediction and feedback checkbox", "Verified expense record"],
        ["Analytics", "SQLite summaries and monthly budgets", "Charts, alerts, remaining budget"],
    ], [1.25, 3.75, 1.9])

    heading(doc, "Data Preprocessing")
    para(doc, "Receipt processing resizes large phone images, crops a document contour when reliable, converts to grayscale, deskews mild rotation, denoises, increases local contrast, enlarges text, and generates an adaptive threshold image. OCR blocks retain text, confidence, and normalized location. Manual descriptions are lowercased, whitespace-normalized, and matched against learned and predefined phrases.")

    heading(doc, "AI Model and Concepts")
    para(doc, "The project is a hybrid AI system. A knowledge base stores deterministic mappings such as autorickshaw to Transportation and user-approved phrases such as sandwich to Snacks. For descriptions without a strong rule, TF-IDF converts unigrams and bigrams into numeric features. Logistic Regression estimates a category probability. Receipt extraction uses confidence and spatial layout, which is a practical form of evidence-based inference.")

    heading(doc, "Algorithms")
    make_table(doc, ["Algorithm or method", "Role in the project"], [
        ["OpenCV preprocessing", "Improves legibility of receipt images."],
        ["EasyOCR", "Recognizes printed receipt text."],
        ["Layout-aware rules", "Pairs final-total labels with nearby amounts and rejects non-final rows."],
        ["TF-IDF", "Represents manual descriptions as weighted text features."],
        ["Logistic Regression", "Predicts a category for unfamiliar descriptions."],
        ["Rule-based expert system", "Applies known and user-learned category decisions."],
    ], [2.2, 4.7])

    heading(doc, "Evaluation Metrics")
    para(doc, "The prototype should be evaluated using field-level exact-match accuracy for merchant and amount, category classification accuracy, precision, recall, macro F1-score, OCR character error rate, and average response time. User correction rate is also useful because it measures when the system requires review. These metrics should be calculated on a labelled hold-out set rather than on training examples.")

    heading(doc, "Dataset and Experimental Setup")
    para(doc, "The project uses a starter labelled expense-description CSV for the category model, manually entered descriptions for rule testing, and a local collection of restaurant and grocery receipt images for OCR review. The application runs locally in Flask with SQLite. The prototype test environment used CPU-based EasyOCR and evaluated clear and degraded receipt photographs separately.")

    heading(doc, "Main Results")
    para(doc, "A clear restaurant receipt was correctly extracted as Bombay Grill House with a final total of 118.46. Manual category rules correctly handled examples such as chips and soft drink, pencils and eraser, autorickshaw, vegetables and milk, petrol, medicine, movie ticket, mobile recharge, and notebook and marker. Faded or logo-heavy receipts can still lose characters or decimal points, so the review screen deliberately asks the user to confirm the result.")

    chart_path = OUT / "functional_results.png"
    draw_bar_chart(chart_path, "Preliminary Functional Validation", ["Manual category examples", "Budget status checks", "Clear receipt extraction"], [100, 100, 100])
    heading(doc, "Quantitative Results Graphs")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(chart_path), width=Inches(6.55))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Figure 1. Functional checks: 9 of 9 manual examples, 2 of 2 budget conditions, and 1 of 1 clear receipt extraction check."), size=8.5, color=MUTED, italic=True)

    heading(doc, "Accuracy Model")
    para(doc, "No overall accuracy percentage is claimed because the available receipt collection does not yet have complete verified merchant, amount, and category labels for a held-out test set. The graph reports only checked functional examples. A rigorous next evaluation should annotate the receipt set, divide it into train, validation, and test partitions, and report field-level exact-match accuracy plus category macro F1-score.")

    heading(doc, "Conclusion")
    para(doc, "The project provides a practical AI-assisted expense tracker that combines OCR, supervised text classification, and rule-based reasoning. Its most useful design choice is human review: users can correct a prediction before saving, and approved manual corrections become private reusable rules. The next stage is a fully labelled receipt benchmark and model comparison to quantify performance across clear, tilted, faded, and multilingual receipts.")

    heading(doc, "Testing and Results")
    make_table(doc, ["Test", "Expected outcome", "Observed result"], [
        ["chips and soft drink", "Snacks", "Snacks"],
        ["pencils and eraser", "Stationery", "Stationery"],
        ["autorickshaw", "Transportation", "Transportation"],
        ["sandwich corrected to Snacks", "Future suggestion uses Snacks", "User-learned rule overrides fallback model"],
        ["Snacks budget 500 and spend 600", "Over-budget alert", "Over by 100"],
        ["Clear Bombay Grill receipt", "Merchant and final total", "Bombay Grill House and 118.46"],
    ], [2.25, 2.35, 2.3])
    doc.save(OUT / "Smart_Expense_AI_Project_Report.docx")


def course_alignment_report():
    doc = Document()
    configure(doc)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Course Alignment for Smart Expense and Receipt Analyzer"), size=22, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Artificial Intelligence Course Policy 2025 26"), size=12, color=MUTED)
    heading(doc, "Alignment Summary")
    para(doc, "The project aligns most directly with the Learning and Expert System units in the supplied Artificial Intelligence policy. It uses supervised text classification for expense categories and rule-based knowledge to map descriptions, merchant names, and user corrections to categories. It also partially aligns with the introductory AI unit through modelling, inference, and practical agent-like decision support. Search, CSP, adversarial search, and unsupervised learning are not central to the current build and should be described as future extensions rather than completed techniques.")

    heading(doc, "Course Topics and Project Mapping")
    make_table(doc, ["Course unit", "Policy topic", "Project alignment"], [
        ["Unit 1", "Modelling, inference, learning, AI applications, agents and environments", "Partial. The expense analyzer models transactions, infers fields and categories, and supports users in a real-world environment."],
        ["Unit 2", "Uninformed and informed search, genetic algorithms, adversarial search", "Not currently implemented. Could be extended with search-based recommendation or optimization."],
        ["Unit 3", "Knowledge representation, logic, inference, forward and backward chaining", "Direct conceptual alignment. Keyword maps, merchant mappings, and learned user rules form a practical knowledge base."],
        ["Unit 4", "Constraint satisfaction problems and backtracking", "Limited alignment. Budget thresholds are constraints, but the project does not currently use CSP or backtracking algorithms."],
        ["Unit 5", "Inductive learning, supervised decision-tree classification, unsupervised K-means", "Direct learning alignment through supervised text classification, although the implemented model is Logistic Regression rather than a decision tree. K-means is not implemented."],
        ["Unit 6", "Expert systems: model, architecture, development and limitations", "Direct alignment. The rule layer and correction memory behave as a small expert system with transparent, editable rules."],
    ], [0.85, 2.2, 3.85])

    heading(doc, "Course Outcome Alignment")
    make_table(doc, ["Course outcome", "Alignment"], [
        ["Explain agents and environment in AI", "Partial: the application can be presented as a decision-support agent operating on receipts, user descriptions, and budget settings."],
        ["Apply heuristic and searching strategies", "Not demonstrated in the present prototype."],
        ["Design knowledge base using expert systems and game playing", "Direct for knowledge base and expert-system aspects; game playing is not relevant."],
        ["Implement supervised and unsupervised learning", "Supervised learning is demonstrated. Unsupervised learning is a recommended extension."],
    ], [2.4, 4.5])

    heading(doc, "Recommended Enhancements for Stronger Syllabus Coverage")
    bullet(doc, "Train and compare a Decision Tree classifier with Logistic Regression on the same labelled expense descriptions.")
    bullet(doc, "Use K-means clustering to group spending patterns by amount, category, or time of month.")
    bullet(doc, "Represent budget planning as a CSP with category limits, target savings, and optional spending constraints.")
    bullet(doc, "Add explainable forward-chaining rules such as IF description contains sandwich THEN category is Snacks.")
    bullet(doc, "Report accuracy, precision, recall, macro F1-score, and OCR field-level exact-match accuracy on a labelled test split.")

    heading(doc, "Conclusion")
    para(doc, "The present project is a credible AI course project because it combines supervised learning with a rule-based expert system and demonstrates inference, feedback learning, and evaluation planning. The strongest academic presentation is to state the implemented methods precisely and position Decision Trees, K-means, and CSP as clearly scoped future work.")
    doc.save(OUT / "AI_Course_Policy_Project_Alignment.docx")


if __name__ == "__main__":
    project_report()
    course_alignment_report()
